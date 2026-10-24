#!/usr/bin/env python3
"""Generate docs/Demo_User_Credentials.docx for O&G Care Dashboard demo logins."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Demo_User_Credentials.docx"

ACCOUNTS = [
    ("ananya.sharma", "Doctor@2025", "Doctor", "Dr. Ananya Sharma — attending obstetrician"),
    ("rahul.mehta", "Resident@2025", "Assistant Doctor", "Dr. Rahul Mehta — resident on labour ward"),
    ("priya.nair", "Nurse@2025", "Nurse", "Priya Nair — labour room charge nurse"),
    ("elena.costa", "Midwife@2025", "Midwife", "Elena Costa — midwife / birthing suite"),
    ("suresh.kumar", "Lift@2025", "Lift Operator", "Suresh Kumar — patient shifting & lift"),
    ("amit.patel", "Wheelchair@2025", "Wheelchair", "Amit Patel — wheelchair transport"),
    ("meena.lab", "Laboratory@2025", "Laboratory", "Meena Rao — lab technician on call"),
    ("deepa.usg", "Ultrasound@2025", "Ultrasound", "Deepa Iyer — ultrasound sonographer"),
]


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    title = doc.add_heading("O&G Care Dashboard — Demo User Credentials", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.add_run(
        "Use these accounts to sign in at the application login screen. "
        "Each username maps to a role-specific dashboard (doctor, nurse, midwife, "
        "lift, wheelchair, laboratory, ultrasound)."
    ).font.size = Pt(11)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run("Important: ").bold = True
    note.add_run(
        "These credentials are for demonstration only. Do not use in production."
    )

    doc.add_heading("Sign-in table", level=1)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, label in enumerate(["Username", "Password", "Role", "Description"]):
        hdr[i].text = label
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True

    for username, password, role, desc in ACCOUNTS:
        row = table.add_row().cells
        row[0].text = username
        row[1].text = password
        row[2].text = role
        row[3].text = desc

    doc.add_paragraph()
    doc.add_heading("Quick reference by role", level=1)
    bullets = [
        "Doctor / Assistant Doctor — patient list, OT scheduling, clinical overview, AI summaries",
        "Nurse / Midwife — labour room checklist, vitals, CTG, assistance calls, ultrasound requests",
        "Lift / Wheelchair — transport and shifting task queue with alarms",
        "Laboratory — incoming lab call queue from labour room",
        "Ultrasound — pending scan requests; upload DICOM and mark reports ready",
    ]
    for item in bullets:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run(f"Generated for project: obgyn-care-dashboard\nPath: {OUT}").font.size = Pt(9)

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
